"""Exporta las tablas de la hoja 'Tablas' del Excel corregido a packages/data/tablas.json.

Las tablas NO se transcriben a mano: este script es la unica via por la que
los datos pasan del Excel (fuente de verdad) a la app y al oraculo.

Uso:  python oracle/export_tables.py "ruta/Darcy Weisbach - corregido.xlsx"
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "packages" / "data" / "tablas.json"

# Factores de caudal: se guardan como fraccion exacta (numerador/denominador)
# para no depender del valor cacheado de una formula de Excel.
FACTORES_CAUDAL = [
    {"unidad": "l/h", "num": 1.0, "den": 3600000.0},
    {"unidad": "l/min", "num": 1.0, "den": 60000.0},
    {"unidad": "l/s", "num": 1.0, "den": 1000.0},
    {"unidad": "m³/h", "num": 1.0, "den": 3600.0},
    {"unidad": "m³/día", "num": 1.0, "den": 86400.0},
    {"unidad": "m³/s", "num": 1.0, "den": 1.0},
    {"unidad": "gpm (US)", "num": 0.003785411784, "den": 60.0},
]


def _rows(ws, first, col_letters):
    r = first
    while ws[f"{col_letters[0]}{r}"].value not in (None, ""):
        yield r, [ws[f"{c}{r}"].value for c in col_letters]
        r += 1


def main(xlsx):
    wb = openpyxl.load_workbook(xlsx)  # valores literales (no formulas) de las tablas
    t = wb["Tablas"]

    materiales = []
    for _, (nombre, kmin, kmax, _kadopt_formula, c_hw, grupo) in _rows(t, 5, "ABCDEF"):
        if not isinstance(kmin, (int, float)):
            break
        materiales.append({"material": nombre, "k_min_mm": float(kmin), "k_max_mm": float(kmax),
                           "c_hw": float(c_hw), "grupo": grupo})

    viscosidad = []
    for _, (temp, nu) in _rows(t, 5, "HI"):
        if not isinstance(temp, (int, float)):
            break
        viscosidad.append({"t_c": float(temp), "nu_m2s": float(nu)})

    def diam(first_row, header_row):
        pns = [float(t[f"{c}{header_row}"].value) for c in "MNOP"]
        out = []
        for _, vals in _rows(t, first_row, "LMNOP"):
            if not isinstance(vals[0], (int, float)):
                break
            for pn, di in zip(pns, vals[1:]):
                if isinstance(di, (int, float)):
                    out.append({"dn": float(vals[0]), "pn": pn, "di_mm": float(di)})
        return out

    singularidades = []
    for _, (nombre, k, nota) in _rows(t, 31, "ABC"):
        if not isinstance(k, (int, float)):
            break
        singularidades.append({"singularidad": nombre, "k": float(k), "nota": nota})

    consts = {t[f"H{r}"].value: float(t[f"I{r}"].value) for r in range(21, 27)}
    constantes = {
        "g": consts["g (gravedad)"],
        "tolerancia": consts["Tolerancia Colebrook |f₁−f₀|"],
        "re_laminar": consts["Re límite laminar"],
        "re_turbulento": consts["Re límite turbulento"],
        "v_min": consts["Velocidad mínima recomendada"],
        "v_max": consts["Velocidad máxima recomendada"],
        "max_iter": 30,
        "f0_liso": 0.02,
        "colebrook_a": 2.51,
        "colebrook_b": 3.71,
    }

    data = {
        "version_tablas": "1.0.0",
        "fuente": Path(xlsx).name,
        "constantes": constantes,
        "materiales": materiales,
        "viscosidad": viscosidad,
        "diametros": {"PVC": diam(5, 4), "PE": diam(28, 27)},
        "singularidades": singularidades,
        "caudal": FACTORES_CAUDAL,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK -> {OUT}  ({len(materiales)} materiales, {len(viscosidad)} temperaturas, "
          f"{len(data['diametros']['PVC'])} DI PVC, {len(data['diametros']['PE'])} DI PE, "
          f"{len(singularidades)} singularidades)")


if __name__ == "__main__":
    main(sys.argv[1])
